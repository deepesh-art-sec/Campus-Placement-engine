"""
Custom Company-Format XLSX & CSV Exporter.
Renders candidate shortlists and interview schedules according to company-specific templates.
"""

from __future__ import annotations

import csv
import io
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.app.api.export.mapping import CompanyExportMapping
from backend.app.domain.models import Student
from backend.app.scheduler.models import ScheduledItem


def export_to_company_xlsx(
    students: list[Student],
    schedule_items: list[ScheduledItem],
    mapping: CompanyExportMapping,
) -> bytes:
    """Exports candidates to an XLSX matching the company's exact layout and column widths."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = mapping.sheet_name or "Candidates"

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    headers = mapping.column_order
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Index schedule by student_id
    sched_map: dict[str, ScheduledItem] = {}
    for item in schedule_items:
        for sid in item.student_ids:
            sched_map[sid] = item

    # Populate candidate rows
    for s in students:
        slot = sched_map.get(s.id)
        row_vals = []
        for h in headers:
            canonical_key = mapping.canonical_mapping.get(h)
            val: Any = ""
            if canonical_key == "roll_no":
                val = s.roll_no
            elif canonical_key == "name":
                val = s.name
            elif canonical_key == "email":
                val = s.email
            elif canonical_key == "branch":
                val = s.branch
            elif canonical_key == "cgpa":
                val = s.cgpa
            elif canonical_key == "tenth_percentage":
                val = s.tenth_percentage
            elif canonical_key == "twelfth_percentage":
                val = s.twelfth_percentage or ""
            elif canonical_key == "round_name":
                val = slot.round_name if slot else ""
            elif canonical_key == "slot_time":
                val = slot.start_time.strftime("%H:%M") if slot else ""
            elif canonical_key == "room_name":
                val = slot.room_name if slot else ""
            elif canonical_key == "panel_name":
                val = slot.panel_name or "" if slot else ""

            row_vals.append(val)

        ws.append(row_vals)

    # Set column widths
    for idx, h in enumerate(headers, start=1):
        col_letter = get_column_letter(idx)
        if h in mapping.column_widths:
            ws.column_dimensions[col_letter].width = mapping.column_widths[h]
        else:
            # Default auto-width
            max_len = max(len(str(h)), 12)
            ws.column_dimensions[col_letter].width = max_len + 3

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_to_company_csv(
    students: list[Student],
    schedule_items: list[ScheduledItem],
    mapping: CompanyExportMapping,
) -> str:
    """Exports candidates to CSV matching the company's exact header ordering."""
    output = io.StringIO()
    writer = csv.writer(output)
    headers = mapping.column_order
    writer.writerow(headers)

    sched_map: dict[str, ScheduledItem] = {}
    for item in schedule_items:
        for sid in item.student_ids:
            sched_map[sid] = item

    for s in students:
        slot = sched_map.get(s.id)
        row_vals = []
        for h in headers:
            canonical_key = mapping.canonical_mapping.get(h)
            val: Any = ""
            if canonical_key == "roll_no":
                val = s.roll_no
            elif canonical_key == "name":
                val = s.name
            elif canonical_key == "email":
                val = s.email
            elif canonical_key == "branch":
                val = s.branch
            elif canonical_key == "cgpa":
                val = s.cgpa
            elif canonical_key == "tenth_percentage":
                val = s.tenth_percentage
            elif canonical_key == "twelfth_percentage":
                val = s.twelfth_percentage or ""
            elif canonical_key == "round_name":
                val = slot.round_name if slot else ""
            elif canonical_key == "slot_time":
                val = slot.start_time.strftime("%H:%M") if slot else ""
            elif canonical_key == "room_name":
                val = slot.room_name if slot else ""
            elif canonical_key == "panel_name":
                val = slot.panel_name or "" if slot else ""

            row_vals.append(val)
        writer.writerow(row_vals)

    return output.getvalue()
