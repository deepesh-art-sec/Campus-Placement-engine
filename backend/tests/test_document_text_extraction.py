"""
Comprehensive Tests for Announcement Text Extraction Feature:
- PyMuPDF (fitz) for PDF digital text & table extraction
- python-docx for DOCX text & table extraction
- Python UTF-8 for TXT extraction
- Scanned PDF detection & Tesseract OCR fallback handling
- Regex/rule-based parsing across all criteria: company, role, branches, CGPA, backlogs, 10th/12th, CTC, dates, location, selection process
- Error handling: unsupported, corrupt, empty, and scanned files
- Coordinator confirmation and SQLite database persistence
"""

from __future__ import annotations

import datetime
import io
import json
import zipfile
import docx
import fitz
import pytest
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.app.api.main import app, seed_default_accounts
from backend.app.extract.document_loader import (
    CorruptDocumentError,
    DocumentLoader,
    EmptyDocumentError,
    ScannedPdfError,
    UnsupportedFormatError,
)
from backend.app.storage.database import SessionLocal
from backend.app.storage.models.entities import (
    CompanyModel,
    DriveModel,
    DriveRequirementVersionModel,
    EventModel,
)


@pytest.fixture
def auth_client():
    seed_default_accounts()
    with TestClient(app) as client:
        res = client.post("/api/auth/login", json={"username": "coordinator", "password": "coord123"})
        assert res.status_code == 200, res.text
        csrf = res.json()["csrf_token"]
        yield client, csrf


def create_synthetic_pdf(text_lines: list[str]) -> bytes:
    """Creates a valid PDF with PyMuPDF text layer."""
    doc = fitz.open()
    page = doc.new_page()
    p = fitz.Point(50, 50)
    for line in text_lines:
        page.insert_text(p, line, fontsize=11)
        p.y += 20
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_synthetic_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    """Creates a valid Word DOCX file using python-docx."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        t = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r_idx, row in enumerate(table_rows):
            for c_idx, val in enumerate(row):
                t.rows[r_idx].cells[c_idx].text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 1. DocumentLoader Unit Tests (PyMuPDF, python-docx, UTF-8, Errors)
# ---------------------------------------------------------------------------


def test_loader_pdf_extraction_with_pymupdf():
    loader = DocumentLoader()
    pdf_bytes = create_synthetic_pdf([
        "Google Campus Recruitment Notice 2026",
        "Role: Software Development Engineer",
        "Eligible branches: Computer Science and Engineering, Information Technology",
        "CTC: 32 LPA",
        "Minimum CGPA required is 8.5 with no active backlogs",
    ])

    text = loader.load_from_bytes(pdf_bytes, "google_notice.pdf")
    assert "Google Campus Recruitment Notice" in text
    assert "Software Development Engineer" in text
    assert "32 LPA" in text


def test_loader_docx_extraction_with_python_docx():
    loader = DocumentLoader()
    docx_bytes = create_synthetic_docx(
        paragraphs=[
            "Microsoft India Placement Circular",
            "Position: Cloud Solutions Architect",
            "Salary: 28.5 LPA",
        ],
        table_rows=[
            ["Criteria", "Requirement"],
            ["Cut-off CGPA", "8.0"],
            ["Active Backlogs", "0"],
        ]
    )

    text = loader.load_from_bytes(docx_bytes, "microsoft.docx")
    assert "Microsoft India Placement Circular" in text
    assert "Cloud Solutions Architect" in text
    assert "Cut-off CGPA | 8.0" in text


def test_loader_txt_extraction_utf8():
    loader = DocumentLoader()
    raw_txt = (
        "Amazon Development Center\n"
        "Job Role: Systems Engineer\n"
        "CTC: 24 LPA\n"
        "Drive Date: 2026-11-20\n"
        "Minimum 75% in 10th & 12th\n"
    )
    txt_bytes = raw_txt.encode("utf-8")

    text = loader.load_from_bytes(txt_bytes, "amazon.txt")
    assert "Amazon Development Center" in text
    assert "Systems Engineer" in text
    assert "24 LPA" in text


def test_loader_rejects_empty_file():
    loader = DocumentLoader()
    with pytest.raises(EmptyDocumentError) as exc:
        loader.load_from_bytes(b"", "empty.txt")
    assert "empty (0 bytes)" in str(exc.value)

    with pytest.raises(EmptyDocumentError) as exc:
        loader.load_from_bytes(b"   \n\n\t  ", "whitespace.txt")
    assert "contains only whitespace" in str(exc.value)


def test_loader_rejects_unsupported_formats():
    loader = DocumentLoader()
    with pytest.raises(UnsupportedFormatError) as exc:
        loader.load_from_bytes(b"some binary executable data", "malware.exe")
    assert "Unsupported document format" in str(exc.value)


def test_loader_rejects_corrupt_pdf():
    loader = DocumentLoader()
    with pytest.raises(CorruptDocumentError) as exc:
        loader.load_from_bytes(b"%PDF-1.4 completely corrupt broken garbage bytes", "broken.pdf")
    assert "Corrupt or invalid PDF file" in str(exc.value)


def test_loader_rejects_corrupt_docx():
    loader = DocumentLoader()
    with pytest.raises(CorruptDocumentError) as exc:
        loader.load_from_bytes(b"PK\x03\x04 not a real docx zip archive", "broken.docx")
    assert "Corrupt or invalid Word document" in str(exc.value)


def test_loader_scanned_pdf_detection():
    loader = DocumentLoader(ocr_enabled=False)
    # Generate PDF with canvas but zero text
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.showPage()
    c.save()
    scanned_bytes = buf.getvalue()

    with pytest.raises(ScannedPdfError) as exc:
        loader.load_from_bytes(scanned_bytes, "scanned_doc.pdf")
    assert "Scanned PDF detected without embedded text layer" in str(exc.value)


# ---------------------------------------------------------------------------
# 2. FastAPI API Endpoint Tests: /api/extract/parse
# ---------------------------------------------------------------------------


def test_api_parse_pdf_all_fields(auth_client):
    client, csrf = auth_client
    pdf_content = create_synthetic_pdf([
        "Cisco Systems Campus Recruitment 2026",
        "Role: Network Software Engineer",
        "Location: Bangalore, India",
        "CTC: 18.5 LPA",
        "Drive Date: 2026-11-28",
        "Mode: On-Campus",
        "Eligibility Criteria:",
        "Minimum CGPA cut-off: 7.5",
        "10th and 12th percentage: minimum 65%",
        "Backlog policy: 0 active backlogs allowed",
        "Eligible branches: Computer Science & Engineering, Electronics & Communication, Mechanical",
        "Education gap: Maximum 1 year education gap",
        "Selection Process:",
        "Round 1: Online Technical Assessment (60 mins)",
        "Round 2: Technical Interview (45 mins)",
        "Round 3: HR Interview (30 mins)",
    ])

    files = {"file": ("cisco_announcement.pdf", pdf_content, "application/pdf")}
    res = client.post("/api/extract/parse", files=files, headers={"X-CSRF-Token": csrf})
    assert res.status_code == 200, res.text
    data = res.json()

    # Verify extracted fields
    assert "Cisco" in data["company_name"]
    assert data["job_role"] == "Network Software Engineer"
    assert data["ctc_lpa"] == 18.5
    assert data["min_cgpa"] == 7.5
    assert data["min_tenth_percentage"] == 65.0
    assert data["min_twelfth_percentage"] == 65.0
    assert data["max_active_backlogs"] == 0
    assert data["max_gap_years"] == 1
    assert data["drive_date"] == "2026-11-28"

    # Verify canonical branch normalization
    assert "CSE" in data["allowed_branches"]
    assert "ECE" in data["allowed_branches"]
    assert "MECH" in data["allowed_branches"]

    # Verify selection process
    rounds = data["selection_process"]
    assert len(rounds) == 3
    assert "Online Technical Assessment" in rounds[0]["name"]
    assert "Technical Interview" in rounds[1]["name"]
    assert "HR Interview" in rounds[2]["name"]

    # Verify field-level provenance
    fields = data["fields"]
    assert fields["cgpa_min"]["extracted_value"] == 7.5
    assert fields["cgpa_min"]["confidence"] >= 0.85
    assert len(fields["cgpa_min"]["source_sentence"]) > 0


def test_api_parse_empty_file_returns_400(auth_client):
    client, csrf = auth_client
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    res = client.post("/api/extract/parse", files=files, headers={"X-CSRF-Token": csrf})
    assert res.status_code == 400
    assert "empty" in res.json()["detail"].lower()


def test_api_parse_corrupt_file_returns_400(auth_client):
    client, csrf = auth_client
    files = {"file": ("corrupt.docx", b"PK\x01\x02 completely random fake bytes", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/extract/parse", files=files, headers={"X-CSRF-Token": csrf})
    assert res.status_code == 400
    assert "corrupt" in res.json()["detail"].lower() or "invalid" in res.json()["detail"].lower()


def test_api_parse_unsupported_format_returns_400(auth_client):
    client, csrf = auth_client
    files = {"file": ("script.exe", b"MZ fake executable binary", "application/octet-stream")}
    res = client.post("/api/extract/parse", files=files, headers={"X-CSRF-Token": csrf})
    assert res.status_code == 400
    assert "unsupported" in res.json()["detail"].lower()


# ---------------------------------------------------------------------------
# 3. Confirmation & SQLite Database Persistence Tests: /api/extract/confirm-save
# ---------------------------------------------------------------------------


def test_api_confirm_save_persists_to_sqlite(auth_client):
    client, csrf = auth_client

    payload = {
        "company_name": "Apex Quantum Dynamics",
        "job_role": "Quantum Systems Engineer",
        "job_location": "Hyderabad",
        "drive_date": "2026-12-05",
        "mode": "on_campus",
        "ctc_lpa": 22.0,
        "min_cgpa": 8.0,
        "min_tenth_percentage": 75.0,
        "min_twelfth_percentage": 70.0,
        "allowed_branches": ["Computer Science and Engineering", "Electronics & Communication", "Mechanical Engineering"],
        "max_active_backlogs": 0,
        "max_gap_years": 1,
        "selection_process": [
            {"round_number": 1, "name": "Online Assessment", "round_type": "oa", "duration_minutes": 60},
            {"round_number": 2, "name": "Technical Interview", "round_type": "technical", "duration_minutes": 45},
        ],
        "confirmed_by": "coordinator",
    }

    res = client.post("/api/extract/confirm-save", json=payload, headers={"X-CSRF-Token": csrf})
    assert res.status_code == 201, res.text
    data = res.json()
    assert data["success"] is True
    drive_id = data["drive_id"]
    assert drive_id.startswith("drv_")

    # Verify normalized branches
    assert "CSE" in data["allowed_branches"]
    assert "ECE" in data["allowed_branches"]
    assert "MECH" in data["allowed_branches"]

    # Verify directly in SQLite database
    db: Session = SessionLocal()
    try:
        # Check DriveModel
        drive = db.execute(select(DriveModel).where(DriveModel.id == drive_id)).scalar_one_or_none()
        assert drive is not None
        assert drive.company_name == "Apex Quantum Dynamics"
        assert drive.ctc_lpa == 22.0
        assert drive.drive_date == datetime.date(2026, 12, 5)
        assert drive.version == 1

        # Check CompanyModel
        company = db.execute(select(CompanyModel).where(CompanyModel.id == drive.company_id)).scalar_one_or_none()
        assert company is not None
        assert company.name == "Apex Quantum Dynamics"

        # Check DriveRequirementVersionModel
        req_v = db.execute(
            select(DriveRequirementVersionModel).where(DriveRequirementVersionModel.drive_id == drive_id)
        ).scalar_one_or_none()
        assert req_v is not None
        assert req_v.version_number == 1
        assert req_v.min_cgpa == 8.0
        assert req_v.min_tenth_percentage == 75.0
        assert req_v.max_active_backlogs == 0
        branches_saved = json.loads(req_v.allowed_branches_json)
        assert "CSE" in branches_saved
        assert "ECE" in branches_saved
        assert "MECH" in branches_saved

        # Check EventModel
        event = db.execute(
            select(EventModel).where(EventModel.aggregate_id == drive_id)
        ).scalar_one_or_none()
        assert event is not None
        assert event.event_type == "ANNOUNCEMENT_EXTRACTION_CONFIRMED"
        payload_logged = json.loads(event.payload_json)
        assert payload_logged["company_name"] == "Apex Quantum Dynamics"
    finally:
        db.close()
